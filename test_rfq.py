import base64
import copy
import io
import json
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
from rfq import HEADERS, decimal_value, extract_pdf, prepare_quote, process, workbook_bytes

ROOT = Path(__file__).parent


class QuoteTests(unittest.TestCase):
    def setUp(self):
        self.cat = [{'code': 'A', 'description': 'Switch', 'description_ar': 'مفتاح', 'unit': 'EA'}]
        self.prices = {'version': 'TEST-1', 'currency': 'AED', 'items': [{'code': 'A', 'unit': 'EA', 'price': '1.00'}]}
        self.row = dict(line='1', code='A', description='Switch', quantity='1', unit='EA', source='Test row')

    def quote(self, **changes):
        return prepare_quote([{**self.row, **changes}], self.cat, self.prices)

    def test_half_up_ties_and_sum_of_rounded_lines(self):
        rows = [{**self.row, 'line': str(i+1), 'quantity': q} for i, q in enumerate(['0.145', '1.005', '2.675'])]
        q = prepare_quote(rows, self.cat, self.prices)
        self.assertEqual([r['amount'] for r in q['rows']], ['0.15', '1.01', '2.68'])
        self.assertEqual(q['complete_total'], '3.84')

    def test_bad_numbers_never_price(self):
        for value in [0.145, True, None, 'NaN', 'Infinity', '1e3', '-1', '0', '1,000', '1.0000001']:
            with self.subTest(value=value):
                q = self.quote(quantity=value)
                self.assertIsNone(q['complete_total'])
                self.assertIsNone(q['rows'][0]['price'])

    def test_unknown_conflict_and_unit_are_reviewed(self):
        for changes in [{'code':'UNKNOWN'}, {'description':'Breaker'}, {'unit':'BOX'}, {'description':''}]:
            with self.subTest(changes=changes):
                q = self.quote(**changes)
                self.assertEqual(q['review_count'], 1)
                self.assertIsNone(q['rows'][0]['amount'])

    def test_missing_price_and_duplicates(self):
        self.prices['items'][0]['code'] = 'OTHER'
        self.assertIn('Missing approved price', self.quote()['rows'][0]['reasons'])
        self.prices['items'][0]['code'] = 'A'
        self.prices['items'].append(copy.deepcopy(self.prices['items'][0]))
        self.assertIn('Duplicate price-list code', self.quote()['rows'][0]['reasons'])
        self.cat.append(copy.deepcopy(self.cat[0]))
        self.assertIn('Duplicate catalogue code', self.quote()['rows'][0]['reasons'])

    def test_price_policy_and_version(self):
        self.prices['items'][0]['price'] = 0.145
        self.assertEqual(self.quote()['review_count'], 1)
        self.prices['currency'] = 'USD'
        with self.assertRaises(ValueError): self.quote()
        self.prices['currency'] = 'AED'
        self.prices['version'] = ''
        with self.assertRaises(ValueError): self.quote()

    def test_formula_like_source_stays_text(self):
        q = self.quote(description='=HYPERLINK("https://example.invalid","bad")')
        book = openpyxl.load_workbook(io.BytesIO(workbook_bytes(q)), data_only=False)
        self.assertEqual(book['Quote EN']['C6'].data_type, 's')
        self.assertEqual(book['عرض السعر AR']['C6'].value, q['rows'][0]['description'])

    def test_source_exception_withholds_complete_total(self):
        q = prepare_quote([self.row], self.cat, self.prices, [{'source':'p1 row3','reason':'Duplicate line','cells':['1']}])
        self.assertIsNone(q['complete_total'])
        self.assertEqual(q['matched_subtotal'], '1.00')
        self.assertEqual(q['review_count'], 1)

    def test_normalization_only(self):
        self.assertEqual(self.quote(code=' a ', description=' SWITCH ', unit='ea')['complete_total'], '1.00')

    def test_real_pdf_to_excel_and_caches(self):
        payload = json.loads(ROOT.joinpath('sample/request.json').read_text(encoding='utf-8'))
        q = process(payload)
        self.assertEqual((len(q['rows']), q['review_count'], q['matched_subtotal'], q['complete_total']), (8,5,'26.16',None))
        data = base64.b64decode(q['xlsx_base64'])
        formula = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        for name in ['Quote EN', 'عرض السعر AR']:
            self.assertEqual(formula[name]['G7'].value, '=ROUND(D7*F7,2)')
            self.assertEqual(cached[name]['G7'].value, .15)
            self.assertEqual(cached[name]['G8'].value, 1.01)
            self.assertEqual(cached[name]['G15'].value, 26.16)
            self.assertEqual(cached[name]['G16'].value, 5)
            self.assertIn(cached[name]['G17'].value, ['', None])
            self.assertIsNone(cached[name]['G9'].value)
        self.assertTrue(formula['عرض السعر AR'].sheet_view.rightToLeft)

    def test_invalid_pdf_and_payload(self):
        for payload in [None, {}, {'pdf_base64':'bad'}, {'pdf_base64':base64.b64encode(b'not a PDF').decode()}]:
            with self.assertRaises(ValueError): process(payload)

    def test_duplicate_source_rows_preserved_as_exceptions(self):
        mockdoc = MagicMock()
        mockdoc.__enter__.return_value.pages = [MagicMock()]
        page = mockdoc.__enter__.return_value.pages[0]
        row = ['1','A','Switch','1','EA']
        page.extract_tables.return_value = [[HEADERS, row, row]]
        with patch('rfq.pdfplumber.open', return_value=mockdoc):
            rows, issues = extract_pdf(b'%PDF-test')
        self.assertEqual((len(rows),len(issues)),(1,1))
        self.assertEqual(issues[0]['cells'], row)

    def test_layout_and_row_limit_fail_closed(self):
        mockdoc = MagicMock()
        mockdoc.__enter__.return_value.pages = [MagicMock()]
        page = mockdoc.__enter__.return_value.pages[0]
        with patch('rfq.pdfplumber.open', return_value=mockdoc):
            for tables in [[], [[['Wrong header']]], [[HEADERS]+[['1','A','Switch','1','EA']]*101]]:
                page.extract_tables.return_value = tables
                with self.assertRaises(ValueError): extract_pdf(b'%PDF-test')


if __name__ == '__main__':
    unittest.main(verbosity=2)
